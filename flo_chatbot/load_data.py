"""FLO bilgi tabanındaki TXT ve DOCX kaynaklarını Qdrant'a yükler."""
import os
import re
import uuid
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
DATA_DIR = BASE_DIR / "data"
COLLECTION_NAME = "flo_bilgi_tabani"
ACTIVE_COLLECTION_NAME = "flo_bilgi_tabani_active"
SUPPORTED_EXTENSIONS = {".txt", ".docx"}
MIN_CHUNK_LENGTH = 80
load_dotenv(BASE_DIR / ".env")

@dataclass
class SourceSection:
    title: str
    text: str
    department: str
    source_file: str

def clean_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()

def department_from_filename(path: Path) -> str:
    name = re.sub(r"^\d+[\s._-]*", "", path.stem)
    name = re.sub(r"\b(departman\s+tanıtımları?|eğitimi|sunumu|notları?)\b", "", name, flags=re.I)
    name = re.sub(r"[\s_-]+", " ", name).strip(" -_")
    return name.title() if name else "Genel"

def docx_sections(path: Path) -> list[tuple[str, str]]:
    document = Document(path)
    sections: list[tuple[str, str]] = []
    title, lines = path.stem, []

    def flush():
        nonlocal lines
        body = clean_text("\n".join(lines))
        if body:
            sections.append((title, body))
        lines = []

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        is_heading = style.startswith(("heading", "başlık")) or bool(re.match(r"^\d{1,3}[.)]\s+\D", text))
        if is_heading:
            flush()
            title = text
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if values:
                lines.append(" | ".join(values))
    flush()
    return sections

def txt_sections(path: Path) -> list[tuple[str, str]]:
    text = clean_text(path.read_text(encoding="utf-8"))
    parts = re.split(r"(?=^(?:#{1,3}\s+|\d{1,3}[.)]\s+\D))", text, flags=re.M)
    return [(part.splitlines()[0].lstrip("# "), part) for part in map(clean_text, parts) if part]

def discover_sources() -> list[Path]:
    if not DATA_DIR.exists():
        return []

    files = [
        p for p in DATA_DIR.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_EXTENSIONS
        and not p.name.startswith("~$")
    ]

    return sorted(files, key=lambda p: p.name.casefold())

def load_sections() -> list[SourceSection]:
    result = []
    for path in discover_sources():
        department = department_from_filename(path)
        parsed = docx_sections(path) if path.suffix.lower() == ".docx" else txt_sections(path)
        for title, body in parsed:
            chunk = clean_text(f"{title}\n{body}")
            if len(chunk) >= MIN_CHUNK_LENGTH:
                section_department = department
                if department == "Genel":
                    section_department = re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", title).strip() or "Genel"
                result.append(SourceSection(title, chunk, section_department, path.name))
    return result

def publish_collection(client, points, vector_size):
    """Eski koleksiyona dokunmadan yükle; yalnızca doğrulanmış veriye geç."""
    from qdrant_client.models import Distance, VectorParams, CreateAlias, CreateAliasOperation, DeleteAlias, DeleteAliasOperation
    name = f"{COLLECTION_NAME}_{uuid.uuid4().hex}"
    client.create_collection(collection_name=name, vectors_config=VectorParams(size=vector_size, distance=Distance.COSINE))
    # Başarısız yükleme yeni koleksiyonda kalır; aktif bilgi kaynağı korunur.
    for offset in range(0, len(points), 128):
        client.upsert(collection_name=name, points=points[offset:offset + 128], wait=True)
    actual = client.count(collection_name=name, exact=True).count
    if actual != len(points):
        raise RuntimeError(f"Yükleme doğrulanamadı: {actual}/{len(points)}. Aktif kaynak değiştirilmedi.")
    aliases = client.get_aliases().aliases
    operations = []
    if any(alias.alias_name == ACTIVE_COLLECTION_NAME for alias in aliases):
        operations.append(DeleteAliasOperation(delete_alias=DeleteAlias(alias_name=ACTIVE_COLLECTION_NAME)))
    operations.append(CreateAliasOperation(create_alias=CreateAlias(collection_name=name, alias_name=ACTIVE_COLLECTION_NAME)))
    client.update_collection_aliases(change_aliases_operations=operations)
    return name


def main() -> None:
    from qdrant_client import QdrantClient
    from qdrant_client.models import PointStruct
    from sentence_transformers import SentenceTransformer
    url, api_key = os.getenv("QDRANT_URL"), os.getenv("QDRANT_API_KEY")
    if not url or not api_key:
        raise RuntimeError("QDRANT_URL ve QDRANT_API_KEY .env dosyasında tanımlanmalıdır.")
    sections = load_sections()
    if not sections:
        raise RuntimeError("Yüklenecek .txt veya .docx bilgi kaynağı bulunamadı.")
    print(f"{len(discover_sources())} kaynakta {len(sections)} bilgi bölümü bulundu.")
    model = SentenceTransformer("sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2")
    vectors = model.encode([s.text for s in sections], show_progress_bar=True, normalize_embeddings=True)
    client = QdrantClient(url=url, api_key=api_key, timeout=30)
    points = {}
    for section, vector in zip(sections, vectors):
        point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{section.source_file}:{section.title}:{section.text}"))
        points[point_id] = PointStruct(id=point_id, vector=vector.tolist(), payload={
            "metin": section.text, "baslik": section.title, "departman": section.department,
            "kaynak": f"FLO GEN - {section.department}", "kaynak_dosya": section.source_file,
        })
    name = publish_collection(client, list(points.values()), model.get_sentence_embedding_dimension())
    print(f"Qdrant yüklemesi tamamlandı: {len(points)} bölüm. Aktif koleksiyon: {name}. Önceki koleksiyonlar korundu.")
    for section in sections:
        print(f"- [{section.department}] {section.title}")

if __name__ == "__main__":
    main()
